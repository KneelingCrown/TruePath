"""Editorial deck: PDF master, fixed-layout PPTX and editable source PPTX.
Requires python-pptx, reportlab, Pillow, pdftoppm and Liberation fonts.
"""
from pathlib import Path
import subprocess
import tempfile
from math import cos,sin,pi
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'presentation'
W,H=1600,900
PAPER='F2F0E7'; INK='11261D'; GREEN='214E3A'; MINT='BDD8B9'; GREY='657468'; RULE='CCD2C6'; DARKRULE='3E5C49'; RED='A85A42'
FONTS=Path('/usr/share/fonts/truetype/liberation')
for n,f in [('sans','LiberationSans-Regular.ttf'),('bold','LiberationSans-Bold.ttf'),('serif','LiberationSerif-Regular.ttf'),('italic','LiberationSerif-Italic.ttf')]:pdfmetrics.registerFont(TTFont(n,str(FONTS/f)))
r=Presentation();r.slide_width=Inches(16);r.slide_height=Inches(9)
r.core_properties.title='True Path — от ответа к пониманию';r.core_properties.author='True Path team'
c=canvas.Canvas(str(OUT/'TruePath.pdf'),pagesize=(W,H));c.setTitle(r.core_properties.title)
notes=[];sl=None;num=0;fg=INK;muted=GREY;rules=RULE

def rgb(v):return RGBColor.from_string(v)
def box(x,y,w,h,fill):
 s=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x/100),Inches(y/100),Inches(w/100),Inches(h/100));s.fill.solid();s.fill.fore_color.rgb=rgb(fill);s.line.fill.background()
 c.setFillColor('#'+fill);c.rect(x,H-y-h,w,h,fill=1,stroke=0)

def dot(x,y,rad=7,col=None):
 col=col or fg
 s=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches((x-rad)/100),Inches((y-rad)/100),Inches(rad*2/100),Inches(rad*2/100));s.fill.solid();s.fill.fore_color.rgb=rgb(col);s.line.fill.background()
 c.setFillColor('#'+col);c.circle(x,H-y,rad,fill=1,stroke=0)

def line(x,y,a,b,col=None,width=1):
 col=col or rules
 s=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x/100),Inches(y/100),Inches(a/100),Inches(b/100));s.line.color.rgb=rgb(col);s.line.width=Pt(width*.72)
 c.setStrokeColor('#'+col);c.setLineWidth(width);c.line(x,H-y,a,H-b)

def t(s,x,y,size=30,col=None,font='sans',maxw=None,lead=1.12):
 col=col or fg;ls=s.split('\n');width=max(pdfmetrics.stringWidth(z,font,size) for z in ls)
 height=(len(ls)-1)*size*lead+size
 assert width<=(maxw or W-x-72)+.1,(num,'text too wide',s,width,maxw)
 assert x>=0 and y>=0 and y+height<=H-35,(num,'text below page',s)
 tb=sl.shapes.add_textbox(Inches(x/100),Inches(y/100),Inches((width+12)/100),Inches((height+12)/100));tf=tb.text_frame;tf.clear();tf.word_wrap=False
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,z in enumerate(ls):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=z;p.font.name='Liberation Serif' if font in ('serif','italic') else 'Liberation Sans';p.font.size=Pt(size*.72);p.font.bold=font=='bold';p.font.italic=font=='italic';p.font.color.rgb=rgb(col);p.line_spacing=Pt(size*lead*.72);p.space_before=p.space_after=Pt(0)
 c.setFont(font,size);c.setFillColor('#'+col);asc=pdfmetrics.getAscent(font)*size/1000
 for i,z in enumerate(ls):c.drawString(x,H-y-asc-i*size*lead,z)

def pic(path,x,y,h):
 from PIL import Image
 iw,ih=Image.open(path).size;w=h*iw/ih
 sl.shapes.add_picture(str(path),Inches(x/100),Inches(y/100),height=Inches(h/100));c.drawImage(str(path),x,H-y-h,width=w,height=h)

def start(section,note,bg=PAPER):
 global sl,num,fg,muted,rules
 if sl is not None:c.showPage()
 num+=1;sl=r.slides.add_slide(r.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(bg)
 c.setFillColor('#'+bg);c.rect(0,0,W,H,fill=1,stroke=0)
 fg=PAPER if bg==INK else INK;muted='A7B9AA' if bg==INK else GREY;rules=DARKRULE if bg==INK else RULE
 # Small branch mark, paired with discreet running heads.
 line(76,58,91,43,fg,2);line(91,43,105,58,fg,2)
 for x,y in [(76,58),(91,43),(105,58)]:dot(x,y,3,fg)
 t('TRUE PATH',124,38,18,font='bold');t(section,1110,40,16,col=muted,maxw=420)
 line(76,815,1524,815)
 t('ОТ ОТВЕТА К ПОНИМАНИЮ',76,838,13,col=muted)
 t(f'{num:02}',1485,833,21,col=fg)
 sl.notes_slide.notes_text_frame.text=note;notes.append((section,note))

start('01  /  Идея','25 секунд. True Path — Android-прототип исследования информации. Мы начинаем с проверяемых оснований и развиваем его в личную систему знаний. Сейчас работаем на учебном потоке; память и адаптация — следующий этап.')
t('True Path',76,164,172,font='serif',maxw=1150)
t('Проверить. Понять. Связать.',84,390,49,col=GREEN,maxw=1170)
t('Из отдельной новости —\nв личную систему знаний.',86,588,32,col=GREY)
# A path, rather than a generic decoration: one material branches into evidence and returns to a conclusion.
for a,b in [((1100,600),(1235,480)),((1235,480),(1390,575)),((1100,600),(1220,718)),((1220,718),(1390,575))]:line(*a,*b,GREEN,2)
for x,y,rr in [(1100,600,10),(1235,480,7),(1220,718,7),(1390,575,14)]:dot(x,y,rr,GREEN)
t('УЧЕБНЫЙ ПРОТОТИП',86,735,15,col=GREY)

start('02  /  Проблема','40 секунд. AI-агент уже умеет исследовать вопрос. Мы не заявляем, что это невозможно повторить промптом. Наша продуктовая гипотеза — связать получение ответа с сохранением оснований и последующим пониманием. Эффект на память ещё нужно измерить.',INK)
t('Ответ есть.',76,160,94,font='serif',col=MINT)
t('Что останется\nчерез неделю?',76,288,94,font='serif',maxw=870,lead=1.02)
line(1050,170,1050,720,width=1)
t('AI уже умеет отвечать.',1110,195,25,col=MINT,maxw=400)
t('Но контекст теряется.\nОснование забывается.\nПохожий случай\nснова кажется новым.',1110,310,29,maxw=400,lead=1.55)
t('Наша задача — сохранить связь между ответом и пониманием.',80,705,29,col=MINT)

start('03  /  Учебный сюжет','40 секунд. Все материалы этого сюжета вымышлены. Новость обещает бесплатный транспорт, но человек хочет знать, можно ли не покупать билет завтра. Нужно проверить область действия, участников и сроки. Недостаточно поставить всей новости одну метку «правда».')
t('Одна фраза.',76,154,83,font='serif')
t('Три разных\nвопроса.',76,263,83,font='serif',col=GREEN,lead=1.02)
t('УЧЕБНАЯ НОВОСТЬ',850,163,15,col=GREY)
t('«В городе запускают\nбесплатный транспорт»',850,209,48,font='serif',maxw=680)
for y,n,head,sub in [(390,'01','Где?','Весь город или один маршрут?'),(520,'02','Для кого?','Все пассажиры или участники пилота?'),(650,'03','Когда?','Постоянно или в дни испытаний?')]:
 line(850,y-20,1524,y-20);t(n,850,y,22,col=GREY);t(head,916,y-6,36,font='serif');t(sub,916,y+48,22,col=GREY,maxw=600)
t('Общий заголовок ещё не даёт основания\nдля личного решения.',80,696,27,maxw=700,col=GREEN)

start('04  /  Основания','45 секунд. Это упрощённая схема учебной проверки, не снимок приложения. Связь ведёт к конкретному условию и документу. Условия пилота нельзя переносить на весь город. В приложении узлы постепенно приходят из подготовленного потока; реальный API пока не подключён.',INK)
t('Вывод, который\nможно разобрать.',76,143,78,font='serif',maxw=1050,lead=1.04)
# Spatial evidence graph: claim above two source branches, one cited condition below.
line(365,400,180,545,MINT,2);line(365,400,570,545,MINT,2)
line(180,545,365,680,MINT,2);line(570,545,365,680,MINT,2)
dot(365,400,11,MINT);dot(180,545,8,MINT);dot(570,545,8,MINT);dot(365,680,13,MINT)
t('Обещание',395,382,22,col=MINT)
t('Публикация',94,573,21);t('Условия пилота',490,573,21);t('Конкретная цитата',395,666,22,col=MINT)
line(790,373,790,727)
t('ОГРАНИЧЕНИЕ',850,387,16,col=MINT)
t('Бесплатно\nв пределах\nпилота.',850,441,77,font='serif',lead=1.01)
t('Это не обещание для всего города.',854,708,27,col=MINT,maxw=660)

start('05  /  Прототип','45 секунд. Показаны реальные Compose-эталоны из репозитория: светлая версия интерфейса. Они не являются фотографиями телефона и не заменяют device-проверку. Учебная лента запускает потоковый граф, а результат сохраняется в истории. В чат также предоставлены более новые тёмные экраны; их оригиналы недоступны в файловом экспорте.')
# Two intact screens are shown as evidence, not as decorative miniatures.
S=ROOT/'app/src/test/snapshots/images'
pic(S/'ru.truepath.app.ui_DesignSnapshotTest_feed.png',100,154,600)
pic(S/'ru.truepath.app.ui_DesignSnapshotTest_graphConclusion.png',410,154,600)
t('Путь от заголовка\nдо основания.',800,238,69,font='serif',maxw=720,lead=1.04)
t('Лента → исследование → граф → история',805,477,27,col=GREEN,maxw=700)
t('Подготовленный поток событий.\nБез выдачи демонстрации\nза реальный фактчекинг.',805,563,29,col=GREY,lead=1.3)
t('COMPOSE / PAPARAZZI · СНИМКИ ИЗ РЕПОЗИТОРИЯ',100,777,13,col=GREY)

start('06  /  Память — следующий этап','50 секунд. Эта функция запланирована, а не реализована. Новое исследование о льготе может напомнить прежнее исследование о транспорте. Общим является способ проверки области действия, не истинность двух выводов. Пользователь видит основания, принимает или отклоняет ассоциацию.',MINT)
t('Следующая новость —\nуже не с нуля.',76,137,82,font='serif',maxw=1400,lead=1.02)
line(245,550,800,550,GREEN,2);line(800,550,1355,550,GREEN,2)
dot(245,550,10,GREEN);dot(800,550,18,GREEN);dot(1355,550,10,GREEN)
t('РАНЬШЕ',110,410,16,col=GREEN);t('Бесплатный\nтранспорт',110,451,42,font='serif')
t('ПРИНЦИП ПРОВЕРКИ',626,616,16,col=GREEN);t('Область действия',626,655,36,font='serif')
t('ТЕПЕРЬ',1120,410,16,col=GREEN);t('Новая льгота\n«для всех»',1120,451,42,font='serif',maxw=410)
t('Объяснимая ассоциация. Не автоматическое подтверждение.',80,754,28,col=GREEN,maxw=1400)

start('07  /  Усвоение — следующий этап','45 секунд. В планируемой системе хранятся основания, личные связи и попытки воспроизведения. Пользователь добровольно отвечает на один вопрос и может вернуться к нему позже. Просмотр или лайк не означают знания. Не обещаем идеальной памяти или доказанного педагогического эффекта.')
t('Сохранить —\nне значит запомнить.',76,140,80,font='serif',maxw=1430,lead=1.04)
t('«Какие условия нужно уточнить,\nпрежде чем менять свои планы?»',80,378,48,font='italic',col=GREEN,maxw=1440)
for x,n,head,body in [(80,'01','Объяснить','Своими словами.'),(590,'02','Вернуться','Когда удобно пользователю.'),(1100,'03','Применить','На другом примере.')]:
 line(x,590,x+420,590);t(n,x,614,18,col=GREY);t(head,x,655,43,font='serif');t(body,x,718,25,col=GREY,maxw=420)

start('08  /  Адаптация — следующий этап','55 секунд. Это два независимых контура. Проверяемые эпизоды источников помогают корректировать маршрут поиска. Интересы и попытки воспроизведения помогают объяснять. Мнение модели или жалоба не являются доказанной ошибкой источника. Журнал — первый этап; автоматические веса только после оценки качества и с откатом.',INK)
t('С опытом — полезнее.',76,145,84,font='serif',col=MINT)
line(800,329,800,680)
t('Лучше\nпроверять.',80,345,70,font='serif',lead=1.02)
t('Понятнее\nобъяснять.',885,345,70,font='serif',lead=1.02)
t('Основания. Независимость.\nИсправления источников.',84,565,29,col=MINT,lead=1.35)
t('Знакомые понятия. Ассоциации.\nРезультаты воспроизведения.',889,565,29,col=MINT,lead=1.35,maxw=640)
t('Предпочтения пользователя не определяют достоверность.',80,742,28,col=MINT)

start('09  /  Фактический статус','40 секунд. Проверяемый код 47f97e4: нами локально выполнены компиляция, 16 unit-тестов и Lint с нулём ошибок и шестью предупреждениями. Это не подтверждает реальный API, жесты на устройстве или усвоение. Справа показаны следующие этапы. Перед защитой обновить показатели по финальной версии.')
t('Что работает сегодня',76,141,80,font='serif')
t('16',76,314,202,font='serif',col=GREEN)
t('unit-тестов пройдено',86,548,37,font='serif')
t('Компиляция успешна\nLint: 0 ошибок / 6 предупреждений',87,625,27,col=GREY,lead=1.4)
line(750,329,750,737)
t('В РАБОЧЕМ ПРОТОТИПЕ',830,341,16,col=GREY)
t('Учебная лента\nПотоковый граф\nЛокальная история',830,387,37,font='serif',lead=1.28)
t('СЛЕДУЮЩИЕ ЭТАПЫ',830,590,16,col=GREY)
t('Реальные API · постоянная память\nАссоциации · усвоение · адаптация',830,635,28,col=GREEN,lead=1.4,maxw=700)
t('ЛОКАЛЬНАЯ ПРОВЕРКА КОДА 47f97e4 · НЕ СТАТУС GITHUB ACTIONS',88,768,13,col=GREY)

start('10  /  Завершение','30 секунд. True Path начинается с проверки информации. Полный продуктовый цикл должен сохранять основания, связывать новые исследования с прежними и поддерживать добровольное усвоение. Ближайшие рубежи — реальная интеграция и локальный сценарий памяти. Готовы к вопросам.',MINT)
t('Знать больше.\nПонимать глубже.',76,174,112,font='serif',maxw=1440,lead=1.03)
t('True Path',84,508,50,font='serif')
t('Каждое исследование помогает следующему.',86,593,34,col=GREEN)
for x,head in [(90,'ПРОВЕРИТЬ'),(640,'ПОНЯТЬ'),(1190,'СВЯЗАТЬ')]:
 dot(x,731,6,GREEN);t(head,x+25,719,18,col=GREEN)
line(250,731,600,731,GREEN);line(790,731,1150,731,GREEN)

start('Запасной  /  Архитектура','Сейчас: Compose, Flow событий и JSON. План M1: Room, постоянные файлы, миграции и экспорт. Обработка на внешних API напрямую из Android при допустимой авторизации. Сервер не разрабатываем; API-ключи не встраиваем в APK. Память локальна, для внешнего вызова берём только нужный контекст.',INK)
t('Память — данные. Не кэш.',76,143,78,font='serif',col=MINT)
for y,n,head,body in [(320,'01','Сейчас','Compose · Flow · JSON-история'),(450,'02','Этап M1','Room · filesDir · миграции · экспорт'),(580,'03','Интеграция','Внешние API · ID задачи · восстановление')]:
 line(80,y-20,1515,y-20);t(n,80,y,22,col=MINT);t(head,170,y-4,37,font='serif');t(body,650,y+3,29,maxw=870)
t('Очистка кэша не удаляет базу. Удаление приложения требует резервной копии.',84,748,23,col=MINT)

start('Запасной  /  Источники','Проверяемый эпизод включает утверждение, основание, дату, тему и исход проверки. Учитываем объём и смещение выборки, независимость источников и исправления. Без доказательства модель не штрафует источник. Автоматические веса остаются этапом M2 после независимого сравнения с базовой политикой.')
t('Каждый эпизод\nможно проверить.',76,145,79,font='serif',maxw=1450,lead=1.05)
t('Утверждение',80,430,32,font='serif');t('Основание',610,430,32,font='serif');t('Пересмотр',1150,430,32,font='serif')
line(80,518,1515,518,GREEN,2)
for x in [80,610,1150]:dot(x,518,7,GREEN)
t('Что было сказано.\nВ какой теме и когда.',80,559,27,col=GREY,lead=1.4)
t('С чем сопоставили.\nКак проверили расхождение.',610,559,27,col=GREY,lead=1.4,maxw=480)
t('Что исправлено.\nЧто нужно перепроверить.',1150,559,27,col=GREY,lead=1.4,maxw=390)
t('Репутация не заменяет доказательство конкретного утверждения.',80,738,30,col=GREEN)

start('Запасной  /  Оценка ценности','Образовательная эффективность пока не измерена. Функциональные проверки отдельно от полезности. Для полезности проверяем принятие и причины отказа от связей, воспроизведение спустя время и перенос на новый пример. Для адаптивной политики — независимые случаи, ошибки, стоимость и задержка. Количество узлов и уверенность модели не показатели истины.',INK)
t('Как проверим\nрезультат.',76,143,83,font='serif',col=MINT,lead=1.04)
for y,n,head,body in [(418,'01','Проверяемость','Есть основания. Сохранены версии.'),(526,'02','Польза','Связь помогает. Знание применяется.'),(634,'03','Качество','Меньше ошибок на независимых случаях.')]:
 line(80,y-19,1515,y-19);t(n,80,y,21,col=MINT);t(head,170,y-4,36,font='serif');t(body,700,y+2,29,maxw=810)

c.save();r.save(str(OUT/'TruePath-editable.pptx'))
# Main PPTX uses verified rendered pages so fonts and layout cannot shift.
fixed=Presentation();fixed.slide_width=Inches(16);fixed.slide_height=Inches(9)
fixed.core_properties.title=r.core_properties.title;fixed.core_properties.author='True Path team'
with tempfile.TemporaryDirectory() as d:
 subprocess.run(['pdftoppm','-scale-to','2400','-png',str(OUT/'TruePath.pdf'),str(Path(d)/'page')],check=True,capture_output=True)
 pages=sorted(Path(d).glob('page-*.png'));assert len(pages)==13
 for p,(_,note) in zip(pages,notes):
  s=fixed.slides.add_slide(fixed.slide_layouts[6]);s.shapes.add_picture(str(p),0,0,width=fixed.slide_width,height=fixed.slide_height);s.notes_slide.notes_text_frame.text=note
fixed.save(str(OUT/'TruePath.pptx'))
(OUT/'SPEAKER_NOTES.md').write_text('# True Path — заметки выступающего\n\nРедакция 2. Десять основных слайдов — около 7 минут. Три запасных — для вопросов.\n\n'+ '\n\n'.join(f'## {i}. {h}\n\n{n}' for i,(h,n) in enumerate(notes,1))+'\n')
print('Built 13 slides: PDF, fixed PPTX, editable PPTX, notes')
